import calendar
from pathlib import Path
from datetime import date
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from services.config import carpeta_fuente

MESES_ES = ["Enero","Febrero","Marzo","Abril","Mayo","Junio",
            "Julio","Agosto","Septiembre","Octubre","Noviembre","Diciembre"]
DIAS_ES  = ["lun","mar","mié","jue","vie","sáb","dom"]

def ruta_informe(año: int, mes: int) -> Path:
    carpeta = carpeta_fuente() / f"{año}_{mes:02d}"
    carpeta.mkdir(parents=True, exist_ok=True)
    return carpeta / f"Mantenimiento_{año}_{mes:02d}.xlsx"

def informe_existe(año: int, mes: int) -> bool:
    return ruta_informe(año, mes).exists()

def generar_excel(año: int, mes: int, dias_por_actividad: dict):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{MESES_ES[mes-1]} {año}"

    _, total_dias = calendar.monthrange(año, mes)

    # --- Estilos base ---
    fill_header = PatternFill("solid", fgColor="1F4E79")
    fill_dia    = PatternFill("solid", fgColor="2E75B6")
    fill_x      = PatternFill("solid", fgColor="C6EFCE")
    font_blanco = Font(color="FFFFFF", bold=True, size=10)
    font_normal = Font(size=10)
    borde = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"),  bottom=Side(style="thin")
    )
    centro = Alignment(horizontal="center", vertical="center")

    # --- Fila 1: Fecha del informe ---
    ws["A1"] = f"Cronograma de Mantenimiento — {MESES_ES[mes-1]} {año}"
    ws["A1"].font = Font(bold=True, size=12)
    ws.merge_cells(f"A1:{_col(10 + total_dias)}1")

    # --- Fila 2: encabezados fijos ---
    fijos = ["N", "Equipo", "Código", "Ubicación", "Actividad",
             "Horas", "Día inicio", "Vencimiento", "Frecuencia", "Alerta"]
    for i, h in enumerate(fijos, 1):
        c = ws.cell(row=2, column=i, value=h)
        c.fill = fill_header
        c.font = font_blanco
        c.alignment = centro
        c.border = borde

    # --- Fila 2 + 3: encabezados de días ---
    for d in range(1, total_dias + 1):
        col = 10 + d
        wd = date(año, mes, d).weekday()
        ws.cell(row=2, column=col, value=DIAS_ES[wd]).fill  = fill_dia
        ws.cell(row=2, column=col).font      = font_blanco
        ws.cell(row=2, column=col).alignment = centro
        ws.cell(row=2, column=col).border    = borde
        ws.cell(row=3, column=col, value=d).fill      = fill_dia
        ws.cell(row=3, column=col).font      = font_blanco
        ws.cell(row=3, column=col).alignment = centro
        ws.cell(row=3, column=col).border    = borde

    # --- Filas de actividades ---
    for fila_idx, (actividad, dias) in enumerate(dias_por_actividad.items(), start=4):
        ws.cell(row=fila_idx, column=1, value=fila_idx - 3).border = borde
        ws.cell(row=fila_idx, column=5, value=actividad).border = borde
        ws.cell(row=fila_idx, column=5).font = font_normal
        for col in range(2, 10):
            ws.cell(row=fila_idx, column=col).border = borde
        for d in dias:
            c = ws.cell(row=fila_idx, column=10 + d, value="X")
            c.fill      = fill_x
            c.alignment = centro
            c.border    = borde

    # --- Anchos de columna ---
    anchos = [5, 12, 10, 12, 28, 7, 10, 12, 12, 20]
    for i, w in enumerate(anchos, 1):
        ws.column_dimensions[_col(i)].width = w
    for d in range(1, total_dias + 1):
        ws.column_dimensions[_col(10 + d)].width = 4

    ruta = ruta_informe(año, mes)
    wb.save(ruta)
    return ruta

def leer_informe(año: int, mes: int) -> dict:
    """Devuelve {nombre_actividad: [lista de días marcados]}"""
    ruta = ruta_informe(año, mes)
    wb = openpyxl.load_workbook(ruta)
    ws = wb.active
    resultado = {}
    _, total_dias = calendar.monthrange(año, mes)
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row):
        nombre = row[4].value  # columna 5 = índice 4
        if not nombre:
            continue
        dias = []
        for d in range(1, total_dias + 1):
            celda = row[10 + d - 1]  # columna 10+d, índice 10+d-1
            if celda.value == "X":
                dias.append(d)
        resultado[nombre] = dias
    return resultado

def _col(n: int) -> str:
    """Número de columna → letra(s) Excel. Ej: 1→A, 27→AA"""
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s